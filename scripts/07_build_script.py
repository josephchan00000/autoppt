#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Stage 7 — 逐字稿 Word（規劃書 §9）。

    - 表格式版面：左欄頁碼 + 投影片主標，右欄逐字稿
    - 依「幕」分節（序幕／每幕的主張句／反方／結語），每節起始插入分頁
    - 頁首寫累計時間軸（「第 12 頁 / 累計 14:30」）
    - 附錄頁另成一節：不計時、不需逐字稿，被問到再翻
    - 文末附「預期提問與備答」一節（從各章 counterpoint 生成）

逐字稿本身由 Claude Code 依 prompts/narration.md 寫進 deck.json 的 narration 欄位，
這支腳本只負責排版成 docx。narration 空白的頁會標【待填】並在最後統計。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import (  # noqa: E402
    DECK_JSON, DIGEST, OUTPUT, SCRIPT_JSON, die, ensure_dirs, info, is_appendix,
    load_project, narration_chars, ok, read_json, safe_filename, step, talk_minutes_range,
    talk_slides, warn, write_json,
)

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

EA_FONT = "微軟正黑體"
FH_RED = RGBColor(0xB8, 0x28, 0x37)
FH_GRAY = RGBColor(0x93, 0x93, 0x96)
FH_DARK = RGBColor(0x26, 0x26, 0x27)


def set_ea(run, name: str = EA_FONT) -> None:
    """python-docx 同樣只設 latin，中文要設 w:eastAsia。"""
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")


def para(container, text: str = "", size: int = 11, bold: bool = False,
         color: RGBColor | None = None, align=None, space_after: int = 6):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color is not None:
            r.font.color.rgb = color
        set_ea(r)
    return p


def mmss(sec: int) -> str:
    return f"{sec // 60}:{sec % 60:02d}"


# ==========================================================================
def main() -> int:
    ap = argparse.ArgumentParser(description="Stage 7：逐字稿 Word")
    ap.add_argument("--deck", default=str(DECK_JSON))
    ap.add_argument("--out", default="")
    args = ap.parse_args()

    ensure_dirs()
    deck_path = Path(args.deck)
    if not deck_path.exists():
        die(f"找不到藍圖：{deck_path}，請先跑 05_outline.py")

    deck = read_json(deck_path)
    meta, slides = deck.get("meta", {}), deck.get("slides", [])
    if not slides:
        die("deck.json 沒有任何 slides")

    cfg = load_project()
    out = Path(args.out) if args.out else OUTPUT / _out_name(meta, cfg)

    step(f"Stage 7 逐字稿（{len(slides)} 頁）")

    doc = Document()
    _setup_styles(doc)
    _cover(doc, meta, slides, cfg)
    stats = _body(doc, meta, slides)
    _qa_section(doc, meta)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    # 同步輸出 06_script.json，供 QA 與改稿使用
    write_json(SCRIPT_JSON, {
        "meta": meta,
        "total_sec": stats["total_sec"],
        "total_chars": stats["chars"],
        "slides": [{"id": s.get("id"), "title": s.get("title"),
                    "role": s.get("role"), "act": s.get("act"),
                    "appendix": is_appendix(s),
                    "duration_sec": s.get("duration_sec", 0),
                    "narration": s.get("narration", ""),
                    "chars": narration_chars(s.get("narration", ""))} for s in slides],
    })

    ok(f"逐字稿 → {out}")
    ok(f"逐字稿索引 → {SCRIPT_JSON}")
    total = stats["total_sec"]
    lo, hi = talk_minutes_range()
    info(f"總時長 {mmss(total)}（目標 {lo}–{hi} 分鐘）／全文 {stats['chars']:,} 字")
    if stats["appendix"]:
        info(f"附錄 {stats['appendix']} 頁不計時，另成一節")
    if stats["empty"]:
        warn(f"{stats['empty']} 頁的 narration 仍空白 —— "
             "請 Claude Code 依 prompts/narration.md 寫進 deck.json 後重跑")
    if not (lo * 60 <= total <= hi * 60):
        warn(f"總時長 {mmss(total)} 不在 {lo}–{hi} 分鐘區間（時間是指引：超時就把次要證據頁"
             "移到附錄，或 python tools/repace_deck.py 重新配速）")
    print()
    info("下一步：python scripts/08_qa.py --all")
    return 0


def _out_name(meta: dict, cfg: dict) -> str:
    title = safe_filename(meta.get("book_title") or cfg["book"].get("title_zh") or "書名", 30)
    return f"FH_{title}_逐字稿.docx"


def _setup_styles(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    for sec in doc.sections:
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)


def _cover(doc: Document, meta: dict, slides: list[dict], cfg: dict) -> None:
    para(doc, "", size=11, space_after=60)
    para(doc, meta.get("book_title", ""), size=26, bold=True, color=FH_RED,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    if meta.get("book_title_en"):
        para(doc, meta["book_title_en"], size=13, color=FH_GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
    para(doc, "分享會逐字稿", size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=40)

    talk = talk_slides(slides)
    total = sum(s.get("duration_sec", 0) for s in talk)
    n_app = len(slides) - len(talk)
    claim = (meta.get("structure") or {}).get("book_claim", "")
    if claim:
        para(doc, claim, size=12, color=FH_DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
    for line in (
        f"講者：{meta.get('dept', '')}　{meta.get('name') or meta.get('presenter', '')}",
        f"日期：{meta.get('date', '')}",
        f"原著：{meta.get('author', '')}",
        f"篇幅：講述 {len(talk)} 頁" + (f"（另附錄 {n_app} 頁）" if n_app else "")
        + f" ／ 預估 {mmss(total)}（目標 {cfg['deck'].get('minutes', 60)} 分鐘，含 Q&A）",
    ):
        para(doc, line, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    doc.add_page_break()


def _section_of(s: dict, act_claims: dict[str, str]) -> str | None:
    """這一頁屬於哪一節。None = 沿用上一節（例如沒掛幕的影片頁）。"""
    if s.get("act") and s["act"] in act_claims:
        return act_claims[s["act"]]
    role = s.get("role")
    if role in ("cover", "summary", "map"):
        return "序幕"
    if role == "counter":
        return "反方"
    if role == "closing":
        return "結語"
    return None


def _slide_row(doc: Document, label: str, s: dict, narration: str) -> None:
    """一頁一列：左欄頁碼＋主標，右欄逐字稿。"""
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    left, right = t.rows[0].cells
    left.width = Cm(4.6)
    right.width = Cm(11.4)

    lp = left.paragraphs[0]
    lp.paragraph_format.space_after = Pt(2)
    r = lp.add_run(label)
    r.font.size = Pt(13)
    r.bold = True
    r.font.color.rgb = FH_RED
    set_ea(r)
    p2 = left.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(s.get("title") or _visual_label(s))
    r2.font.size = Pt(10)
    r2.font.color.rgb = FH_DARK
    set_ea(r2)
    if s.get("subtitle"):
        p3 = left.add_paragraph()
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(s["subtitle"])
        r3.font.size = Pt(9)
        r3.font.color.rgb = FH_GRAY
        set_ea(r3)

    rp = right.paragraphs[0]
    rp.paragraph_format.space_after = Pt(0)
    rr = rp.add_run(narration)
    rr.font.size = Pt(11)
    set_ea(rr)
    para(doc, "", size=8, space_after=10)


def _visual_label(s: dict) -> str:
    """引言頁／大數字頁沒有主標，用內容當左欄標籤。"""
    if s.get("quote"):
        return "引言：" + (s["quote"].get("zh") or s["quote"].get("text") or "")[:24]
    if s.get("stat"):
        return "大數字：" + str(s["stat"].get("value", "")) + " " + (s["stat"].get("label") or "")[:16]
    return "（無主標）"


def _body(doc: Document, meta: dict, slides: list[dict]) -> dict:
    act_claims = {a["id"]: a["claim"] for a in (meta.get("structure") or {}).get("acts") or []}
    talk = talk_slides(slides)
    appendix = [s for s in slides if is_appendix(s)]
    cum = empty = chars = 0
    cur_section: str | None = None

    for i, s in enumerate(talk, start=1):
        section = _section_of(s, act_claims)
        if section and section != cur_section:
            if cur_section is not None:
                doc.add_page_break()
            cur_section = section
            para(doc, section, size=16, bold=True, color=FH_RED, space_after=10)

        dur = int(s.get("duration_sec", 0))
        narration = (s.get("narration") or "").strip()
        if not narration:
            narration = "【待填】依 prompts/narration.md 寫入 deck.json 的 narration 欄位"
            empty += 1
        else:
            chars += narration_chars(narration)

        para(doc, f"第 {i} 頁 / 累計 {mmss(cum)} ／ 本頁 {dur} 秒", size=9, color=FH_GRAY, space_after=2)
        _slide_row(doc, f"P{i}", s, narration)
        cum += dur

    if appendix:
        doc.add_page_break()
        para(doc, "附錄：備用頁", size=16, bold=True, color=FH_RED, space_after=4)
        para(doc, "不計入時間、不需逐字稿。Q&A 被問到相關題目時再翻，左欄是頁碼。",
             size=10, color=FH_GRAY, space_after=10)
        for j, s in enumerate(appendix, start=len(talk) + 1):
            narration = (s.get("narration") or "").strip() or "（備用頁，無逐字稿）"
            _slide_row(doc, f"P{j}", s, narration)

    return {"total_sec": cum, "empty": empty, "chars": chars,
            "talk": len(talk), "appendix": len(appendix)}


def _qa_section(doc: Document, meta: dict) -> None:
    """文末附「預期提問與備答」一節（從各章 counterpoint 生成）。"""
    doc.add_page_break()
    para(doc, "預期提問與備答", size=18, bold=True, color=FH_RED, space_after=6)
    para(doc, "以下問題由各章的 counterpoint（最值得挑戰的地方）生成，"
              "請在會前補上你自己的答法。", size=10, color=FH_GRAY, space_after=14)

    rows = []
    for c in meta.get("chapters", []):
        p = DIGEST / f"{c['ch_id']}.json"
        if not p.exists():
            continue
        d = read_json(p)
        cp = (d.get("counterpoint") or "").strip()
        if cp:
            rows.append((c["title"], cp))

    if not rows:
        para(doc, "（找不到 work/03_digest/*.json，無法生成。跑完 Stage 3 後重跑本階段。）",
             size=11, color=FH_GRAY)
        return

    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for cell, label, w in zip(hdr, ("章節", "可能被問", "我的答法"), (3.6, 6.4, 6.0)):
        cell.width = Cm(w)
        pr = cell.paragraphs[0]
        r = pr.add_run(label)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = FH_DARK
        set_ea(r)

    for title, cp in rows:
        cells = t.add_row().cells
        for cell, text, size, w in zip(
            cells,
            (title, f"「{cp}」這點你怎麼看？", ""),
            (10, 10, 10),
            (3.6, 6.4, 6.0),
        ):
            cell.width = Cm(w)
            pr = cell.paragraphs[0]
            pr.paragraph_format.space_after = Pt(2)
            r = pr.add_run(text)
            r.font.size = Pt(size)
            set_ea(r)

    para(doc, "", space_after=12)
    para(doc, "時間控制提醒", size=13, bold=True, color=FH_RED, space_after=4)
    lo, hi = talk_minutes_range()
    for line in (
        f"・講述目標 {lo}–{hi} 分鐘，留 5 分鐘 Q&A。時間是指引，不是門檻。",
        "・落後時跳過證據頁的細節，不要跳過每一幕的主張頁與意涵頁——那是論證的骨架。",
        "・主張頁籤只講一句話（20 秒內）；引言頁、大數字頁不停留。",
        "・附錄頁不講，被問到再翻。",
    ):
        para(doc, line, size=10, space_after=3)


if __name__ == "__main__":
    sys.exit(main())
