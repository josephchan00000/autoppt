#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Stage 8 — 自動品管（不可略過）。輸出 output/qa_report.md。

任何一項 FAIL 就不准交付：
    敘事結構  導讀體：有作者頁、執行摘要在序幕；每一主線章有章名頁籤（官方章名或原文）／
              大綱頁／示意圖或今天的數字；有全書意涵、反方、結語；無殘留【待填】。
              節奏（頁籤秒數、序幕佔比、單頁停留）只 WARN
    原書用字  每一主線章的投影片至少出現該章一個作者用語（digest.key_terms）；否則 WARN
    故事      每一主線章的逐字稿要講到該章選的故事（大綱頁 story_hint 的人物）；否則 WARN
    條列樣式  文字頁必為 chain / labeled / prose 之一（裸條列 FAIL）；標籤 ≤5 字
    溢排      每頁內文估算行數 ≤ 8；主標 ≤ 14 字；副標 ≤ 21 字
    資料來源  每一頁（封面／頁籤／結尾除外）都有 sources，且非空字串
    外部連結  所有 url HTTP 200；死連結列出
    具體性    每頁 body 至少含一個數字／年份／專有名詞；否則標 WARN
    版型純度  每張 slide 的 layout name 必須在模板 11 種之內；無自建 textbox（資料來源行除外）
    字型      掃描所有 run，a:ea typeface 必須是 微軟正黑體
    逐字稿    每頁都有 narration（附錄除外）；字數與總時長只 WARN——時間是指引，品質才是門檻
    節奏      連續條列頁 ≤ 2；視覺頁佔比 ≥ 40%
    對岸用語  黑名單掃描
    頁數      參考區間，只 WARN
    視覺      LibreOffice 轉 PNG 全頁截圖，輸出到 output/preview/

用法：
    python scripts/08_qa.py --all             全部檢查 + 產報告
    python scripts/08_qa.py --check-env       只檢查環境與相依
    python scripts/08_qa.py --check-sources   只做 URL HTTP 全檢（Stage 4 後用）
    python scripts/08_qa.py --check-deck      只檢查 deck.json 文案規則（Stage 5 後用）
    python scripts/08_qa.py --check-narration 只檢查逐字稿字數
    python scripts/08_qa.py --preview         只產預覽圖
"""
from __future__ import annotations

import argparse
import re
import subprocess
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import (  # noqa: E402
    BODY_STYLES, CARD_COMFORT_PT, DECK_JSON, DIGEST, EVIDENCE, OUTPUT, TEMPLATE, _c,
    body_item_text, card_fit, card_per_page, ensure_dirs,
    estimate_lines, format_timecode, has_concrete, have_cmd, info, is_appendix, limits,
    load_project, load_spec, narration_chars, ok, read_json, soffice_bin, step,
    talk_minutes_range, talk_slides, target_slides, visual_len, warn,
)

ALLOWED_CUSTOM_SHAPES = {
    "SourceLine", "ChartTitle", "VideoTitle", "VideoInfo",
    # 版面 9 的視覺頁元件（06_build_pptx.py 畫的）
    "VisualTitle", "VisualText", "QuoteBar", "QuoteText", "QuoteAttrib",
    "StatValue", "StatLabel", "TimelineAxis", "TimelineDot", "TimelineLabel",
    "FlowBox", "FlowArrow", "FlowCaption", "FlowLoopLabel",
    "SplitBox", "SplitHeader", "SplitBody", "SplitCaption", "DividerNote",
    # 卡片式文字頁（06_build_pptx.draw_cards）
    "CardBox", "CardChip", "CardNum", "CardLink", "CardText", "ProseBar", "ProseText",
    "ImageSlot",
}
EA_EXPECT = "微軟正黑體"


class Result:
    """一項檢查的結果。"""

    def __init__(self, name: str, criterion: str):
        self.name = name
        self.criterion = criterion
        self.status = "PASS"     # PASS / WARN / FAIL / SKIP
        self.details: list[str] = []

    def fail(self, msg: str) -> None:
        self.status = "FAIL"
        self.details.append(msg)

    def warn(self, msg: str) -> None:
        if self.status == "PASS":
            self.status = "WARN"
        self.details.append(msg)

    def skip(self, msg: str) -> None:
        self.status = "SKIP"
        self.details.append(msg)

    def note(self, msg: str) -> None:
        self.details.append(msg)

    def print_line(self) -> None:
        color = {"PASS": "32", "WARN": "33", "FAIL": "31", "SKIP": "90"}[self.status]
        mark = {"PASS": "✓", "WARN": "!", "FAIL": "✗", "SKIP": "–"}[self.status]
        print(_c(color, f"  {mark} {self.status:<4} {self.name}"))
        for d in self.details[:12]:
            print(f"        {d}")
        if len(self.details) > 12:
            print(f"        …（另有 {len(self.details) - 12} 項，詳見 qa_report.md）")


# ==========================================================================
def main() -> int:
    ap = argparse.ArgumentParser(description="Stage 8：自動品管")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--check-env", action="store_true")
    ap.add_argument("--check-sources", action="store_true")
    ap.add_argument("--check-deck", action="store_true")
    ap.add_argument("--check-narration", action="store_true")
    ap.add_argument("--preview", action="store_true")
    ap.add_argument("--pptx", default="")
    ap.add_argument("--timeout", type=float, default=12.0, help="URL 檢查逾時秒數")
    args = ap.parse_args()

    ensure_dirs()
    if not any([args.all, args.check_env, args.check_sources, args.check_deck,
                args.check_narration, args.preview]):
        ap.print_help()
        return 2

    if args.check_env:
        return check_env()

    results: list[Result] = []

    if args.check_sources:
        step("URL 連結檢查")
        r = check_links(collect_urls_from_evidence(), args.timeout)
        r.print_line()
        return 0 if r.status in ("PASS", "WARN", "SKIP") else 1

    deck = read_json(DECK_JSON) if DECK_JSON.exists() else None

    if args.check_deck or args.check_narration:
        if deck is None:
            print(_c("31", "  ✗ 找不到 work/05_deck.json，請先跑 05_outline.py"))
            return 1
        step("deck.json 檢查")
        if args.check_deck:
            results += [check_structure(deck), check_wording(deck), check_stories(deck),
                        check_body_styles(deck), check_overflow(deck), check_sources(deck),
                        check_concrete(deck), check_rhythm(deck), check_ai_tone(deck),
                        check_banned_terms(deck), check_page_count(deck)]
        if args.check_narration:
            results.append(check_narration(deck))
        for r in results:
            r.print_line()
        return 0 if not any(r.status == "FAIL" for r in results) else 1

    pptx = Path(args.pptx) if args.pptx else find_pptx()

    if args.preview:
        r = make_preview(pptx)
        r.print_line()
        return 0

    # ---- --all ----
    step("Stage 8 自動品管")
    if deck is None:
        print(_c("31", "  ✗ 找不到 work/05_deck.json，請先跑 05_outline.py"))
        return 1

    results.append(check_structure(deck))
    results.append(check_wording(deck))
    results.append(check_stories(deck))
    results.append(check_body_styles(deck))
    results.append(check_overflow(deck))
    results.append(check_sources(deck))
    results.append(check_links(collect_urls_from_deck(deck), args.timeout))
    results.append(check_concrete(deck))
    results.append(check_rhythm(deck))
    results.append(check_layout_purity(pptx))
    results.append(check_fonts(pptx))
    results.append(check_narration(deck))
    results.append(check_ai_tone(deck))
    results.append(check_banned_terms(deck))
    results.append(check_page_count(deck))
    results.append(check_pdf(pptx))
    results.append(make_preview(pptx))

    for r in results:
        r.print_line()

    report = write_report(results, deck, pptx)
    print()
    n_fail = sum(1 for r in results if r.status == "FAIL")
    n_warn = sum(1 for r in results if r.status == "WARN")
    n_skip = sum(1 for r in results if r.status == "SKIP")
    ok(f"報告 → {report}")
    if n_fail:
        print(_c("31", f"\n  ✗ {n_fail} 項 FAIL —— 不准交付。修完 deck.json 後跑 make revise"))
        return 1
    print(_c("32", f"\n  ✓ 全部通過（{n_warn} 項 WARN，{n_skip} 項 SKIP）"))
    if n_warn:
        info("WARN 不擋交付，但建議看一下 qa_report.md")
    return 0


# ==========================================================================
def find_pptx() -> Path | None:
    cands = sorted(OUTPUT.glob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return cands[0] if cands else None


def slides_of(deck: dict) -> list[dict]:
    return deck.get("slides", [])


# --- 頁數（參考值，只 WARN）--------------------------------------------------
def check_page_count(deck: dict) -> Result:
    """頁數只是參考：品質優先，不裁頁。講得完不完看「逐字稿」那一項的總時長。"""
    lo, hi = target_slides()
    talk = talk_slides(slides_of(deck))
    n_app = len(slides_of(deck)) - len(talk)
    n = len(talk)
    r = Result("頁數", f"講述頁 {lo}–{hi} 為參考區間（只 WARN，不裁頁）")
    tail = f"（另有附錄 {n_app} 頁不計）" if n_app else ""
    if lo <= n <= hi:
        r.note(f"{n} 頁講述頁{tail}")
    else:
        r.warn((f"講述頁只有 {n} 頁，低於參考下限 {lo}" if n < lo
                else f"講述頁 {n} 頁，高於參考上限 {hi}") + tail)
        r.note("視覺頁多的簡報本來就頁數多、每頁短；看逐字稿的總時長才準")
    return r


# --- 2. 溢排 --------------------------------------------------------------
def check_overflow(deck: dict) -> Result:
    lim = limits()
    max_lines = int(lim.get("max_body_lines", 8))
    t_max = int(lim.get("title_max_chars", 14))
    # 圖表頁走版面 9，主標是 32pt textbox（寬 8,208,144 EMU），比內頁1-1 的
    # 36pt 主標框（寬 6,840,760 EMU）放得下更多字，所以另設一個上限。
    ct_max = int(lim.get("chart_title_max_chars", 20))
    s_max = int(lim.get("subtitle_max_chars", 21))
    # 視覺頁（版面 9）的副標是 18pt textbox，比內頁的 24pt 副標框寬得多
    cs_max = int(lim.get("chart_subtitle_max_chars", 30))
    b_max = int(lim.get("body_max_chars_per_slide", 160))
    l1_max = int(lim.get("bullet_l1_max_chars", 40))
    max_lv = int(lim.get("max_outline_level", 1))

    r = Result("溢排", f"卡片頁看最小字級塞不塞得下（縮到 {CARD_COMFORT_PT}pt 以下 WARN）；"
                       f"其餘內文 ≤{max_lines} 行；主標 ≤{t_max} 字"
                       f"（空白內頁的視覺頁 ≤{ct_max}）；副標 ≤{s_max} 字")
    for s in slides_of(deck):
        sid, kind = s.get("id"), s.get("kind")
        title = s.get("title") or ""
        sub = s.get("subtitle") or ""
        body = s.get("body") or []

        # image / table / split 走內頁2-1，主標是母片 36pt 的 placeholder，
        # 容量與內容頁相同；其餘視覺頁走空白內頁的自建 textbox，放得下比較多字。
        on6 = any(s.get(k) for k in ("image", "table", "split"))
        if kind not in ("cover", "divider", "closing", "toc", "video"):
            cap = ct_max if (kind == "chart" and not on6) else t_max
            if visual_len(title) > cap:
                r.fail(f"{sid} 主標 {visual_len(title):.0f} 字 > {cap}：「{title[:24]}」")
        if sub:
            scap = cs_max if (kind in ("chart", "cover") and not on6) else s_max
            if visual_len(sub) > scap:
                r.fail(f"{sid} 副標 {visual_len(sub):.0f} 字 > {scap}：「{sub[:30]}」")

        style = s.get("style") if s.get("style") in BODY_STYLES else None
        if kind == "divider" and visual_len(title) > t_max:
            r.warn(f"{sid} 頁籤 {visual_len(title):.0f} 字 > {t_max}，44pt 會折成三行")

        if style and kind == "content":
            # 卡片式：字級由版面引擎自己縮，所以問的不是「幾行」而是
            # 「最小字級還塞不塞得下」「縮到幾點還讀得舒服」。與引擎同一份算法。
            fit = card_fit(body[:card_per_page(style)], style)
            if not fit["fits"]:
                r.fail(f"{sid} 卡片連 {fit['size_pt']}pt 都塞不下，要砍字或拆頁")
            elif fit["size_pt"] < CARD_COMFORT_PT:
                r.warn(f"{sid} 卡片縮到 {fit['size_pt']}pt（舒服的下限 {CARD_COMFORT_PT}pt），"
                       "字再短一點會好看很多")
        else:
            lines = estimate_lines(body, 24, style)
            if lines > max_lines:
                r.fail(f"{sid} 內文估算 {lines} 行 > {max_lines}（24pt）")

        total = sum(visual_len(body_item_text(style, b, i)) for i, b in enumerate(body))
        if total > b_max:
            r.fail(f"{sid} 內文總字數 {total:.0f} > {b_max}")

        for b in body:
            lv = int(b.get("level", 0))
            # 封面與章序地圖（toc）的階層是母片固定版面（分部 → 章名），不算條列的第二層
            if lv > max_lv and kind not in ("cover", "toc"):
                r.fail(f"{sid} 出現第 {lv + 1} 層（最多 {max_lv + 1} 層）")
            # prose 的段落上限另外在「條列樣式」檢查
            if lv == 0 and style != "prose" and visual_len(b.get("text") or "") > l1_max:
                r.fail(f"{sid} 要點 {visual_len(b.get('text') or ''):.0f} 字 > {l1_max}")

        if kind == "content" and style != "prose":
            l1 = [b for b in body if int(b.get("level", 0)) == 0]
            blo, bhi = lim.get("bullets_l1_range", [2, 3])
            # chain 是大綱頁，prompts/guide.md 寫的就是「大綱 2–4 步」，
            # 上限跟著版面引擎一頁放得下的張數走，不要拿 labeled 的門檻去卡它
            if style == "chain":
                bhi = max(bhi, card_per_page(style))
            if l1 and not (blo <= len(l1) <= bhi):
                r.warn(f"{sid} 要點 {len(l1)} 條，建議 {blo}–{bhi} 條")
    if r.status == "PASS":
        r.note(f"{len(slides_of(deck))} 頁全部在版面容量內")
    return r


# --- 3. 資料來源 ----------------------------------------------------------
def check_sources(deck: dict) -> Result:
    r = Result("資料來源", "每頁（封面／頁籤／結尾除外）都有 sources 且非空字串")
    exempt = {"cover", "divider", "closing", "wish"}
    missing = []
    for s in slides_of(deck):
        if s.get("kind") in exempt:
            continue
        srcs = s.get("sources") or []
        labels = [(x.get("label") or "").strip() for x in srcs if isinstance(x, dict)]
        labels = [x for x in labels if x]
        if not labels:
            missing.append(s.get("id"))
        elif any("【待填" in x for x in labels):
            r.warn(f"{s.get('id')} 來源含【待填】標記")
    if missing:
        r.fail(f"{len(missing)} 頁沒有資料來源：{', '.join(map(str, missing[:15]))}"
               + (" …" if len(missing) > 15 else ""))
    else:
        r.note("所有需要來源的頁面都有")
    return r


# --- 4. 外部連結 ----------------------------------------------------------
def collect_urls_from_deck(deck: dict) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for s in slides_of(deck):
        for src in s.get("sources") or []:
            if isinstance(src, dict) and src.get("url"):
                out.setdefault(src["url"], []).append(str(s.get("id")))
    return out


def collect_urls_from_evidence() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for p in sorted(EVIDENCE.glob("ch*.json")):
        try:
            d = read_json(p)
        except Exception:  # noqa: BLE001
            continue
        for key in ("verified", "taiwan_lens", "extensions", "chart_candidates"):
            for item in d.get(key) or []:
                if isinstance(item, dict) and item.get("source_url"):
                    out.setdefault(item["source_url"], []).append(f"{p.stem}.{key}")
    return out


def check_links(urls: dict[str, list[str]], timeout: float) -> Result:
    r = Result("外部連結", "所有 url HTTP 200；死連結列出")
    if not urls:
        r.note("沒有外部連結需要檢查")
        return r
    try:
        import requests
    except ImportError:
        r.skip("未安裝 requests，無法檢查連結（pip install requests）")
        return r

    info(f"檢查 {len(urls)} 個連結…")
    headers = {"User-Agent": "Mozilla/5.0 (compatible; FH-deck-QA/1.0)"}

    def probe(url: str) -> tuple[str, int | str]:
        try:
            resp = requests.head(url, timeout=timeout, allow_redirects=True, headers=headers)
            if resp.status_code >= 400:
                resp = requests.get(url, timeout=timeout, allow_redirects=True,
                                    headers=headers, stream=True)
            return url, resp.status_code
        except Exception as e:  # noqa: BLE001
            return url, type(e).__name__

    with ThreadPoolExecutor(max_workers=8) as ex:
        results = list(ex.map(probe, urls.keys()))

    dead, okc, netfail = [], 0, 0
    NET_ERRORS = {"ProxyError", "ConnectionError", "ConnectTimeout", "ReadTimeout",
                  "Timeout", "SSLError", "TooManyRedirects", "ChunkedEncodingError"}
    for url, code in results:
        where = ", ".join(urls[url][:3])
        if code == 200:
            okc += 1
        elif isinstance(code, int) and code in (301, 302, 303, 307, 308, 403, 405):
            r.warn(f"HTTP {code}（可能仍有效，請人工確認）{url}  ← {where}")
        else:
            if code in NET_ERRORS:
                netfail += 1
            dead.append((url, code, where))

    # 全部都是連線層失敗 → 是這台機器連不出去，不是連結真的死掉。
    # 判 SKIP 而不是 FAIL，否則在受限網路環境會永遠卡住交付。
    if netfail and netfail == len(urls):
        r.skip(f"{len(urls)} 個連結全部連線失敗（{results[0][1]}），"
               "判定為本機無法對外連線而非死連結。")
        r.note("請在可正常上網的環境重跑：python scripts/08_qa.py --check-sources")
        for url, code, where in dead[:10]:
            r.note(f"未驗證：{url}  ← {where}")
        return r

    if dead:
        for url, code, where in dead:
            r.fail(f"死連結 HTTP {code}：{url}  ← {where}")
    r.note(f"{okc}/{len(urls)} 個連結回應 200"
           + (f"，{netfail} 個連線失敗" if netfail else ""))
    return r


# --- 5. 具體性 ------------------------------------------------------------
def check_concrete(deck: dict) -> Result:
    r = Result("具體性", "每頁 body 至少含一個數字／年份／專有名詞，否則 WARN")
    exempt = {"cover", "divider", "closing", "toc", "chart", "video"}
    for s in slides_of(deck):
        if s.get("kind") in exempt:
            continue
        body = s.get("body") or []
        text = " ".join((b.get("text") or "") for b in body)
        if not text.strip():
            continue
        if not has_concrete(text) and not has_concrete(s.get("title") or ""):
            r.warn(f"{s.get('id')} 全是抽象概念，沒有數字／年份／案例名："
                   f"「{(s.get('title') or '')[:20]}」")
    if r.status == "PASS":
        r.note("每一頁都有具體物")
    return r


# --- 6. 版型純度 ----------------------------------------------------------
def check_layout_purity(pptx: Path | None) -> Result:
    spec = load_spec()
    allowed = {L["name"] for L in spec["layouts"]}
    r = Result("版型純度", f"layout 必須在模板 {len(allowed)} 種之內；無自建 textbox")
    if pptx is None or not pptx.exists():
        r.skip("找不到 output/*.pptx，請先跑 06_build_pptx.py")
        return r
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        r.skip("未安裝 python-pptx")
        return r

    prs = Presentation(str(pptx))
    used = Counter()
    for i, sl in enumerate(prs.slides, start=1):
        name = sl.slide_layout.name
        used[name] += 1
        if name not in allowed:
            r.fail(f"第 {i} 頁用了模板外的版面「{name}」")
        for shp in sl.shapes:
            if shp.is_placeholder:
                continue
            if shp.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shp.name not in ALLOWED_CUSTOM_SHAPES:
                r.fail(f"第 {i} 頁有自建 textbox「{shp.name}」"
                       f"（只允許 {sorted(ALLOWED_CUSTOM_SHAPES)}）")
            if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                r.fail(f"第 {i} 頁殘留裝飾群組 shape，放圖表前應移除")
    r.note("版面使用：" + "、".join(f"{k}×{v}" for k, v in used.most_common()))
    return r


# --- 7. 字型 --------------------------------------------------------------
def check_fonts(pptx: Path | None) -> Result:
    r = Result("字型", f"所有 run 的 a:ea typeface = {EA_EXPECT}")
    if pptx is None or not pptx.exists():
        r.skip("找不到 output/*.pptx")
        return r
    try:
        from pptx import Presentation
        from pptx.oxml.ns import qn
    except ImportError:
        r.skip("未安裝 python-pptx")
        return r

    prs = Presentation(str(pptx))
    total = bad = 0
    examples: list[str] = []
    for i, sl in enumerate(prs.slides, start=1):
        frames = [shp.text_frame for shp in sl.shapes if shp.has_text_frame]
        if sl.has_notes_slide:
            frames.append(sl.notes_slide.notes_text_frame)
        for tf in frames:
            for para in tf.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    total += 1
                    rPr = run._r.find(qn("a:rPr"))
                    ea = rPr.find(qn("a:ea")) if rPr is not None else None
                    if ea is None or ea.get("typeface") != EA_EXPECT:
                        bad += 1
                        if len(examples) < 8:
                            got = ea.get("typeface") if ea is not None else "（未設定）"
                            examples.append(f"第 {i} 頁「{run.text[:18]}」ea={got}")
    if bad:
        r.fail(f"{bad}/{total} 個 run 的 a:ea 不是 {EA_EXPECT}（換機器開會變新細明體）")
        for e in examples:
            r.note(e)
    else:
        r.note(f"{total} 個 run 全數正確")
    return r


# --- 8. 逐字稿 ------------------------------------------------------------
def check_narration(deck: dict) -> Result:
    cfg = load_project().get("narration", {})
    cpm = int(cfg.get("chars_per_minute", 220))
    tol = float(cfg.get("tolerance", 0.25))
    lo_m, hi_m = talk_minutes_range()

    r = Result("逐字稿", f"每頁都有 narration（附錄除外）；字數 = duration_sec × {cpm}/60 "
                        f"±{tol:.0%} 與總時長 {lo_m}–{hi_m} 分鐘只 WARN")
    slides = talk_slides(slides_of(deck))
    empty = []
    total_sec = 0
    video_sec = 0
    for s in slides:
        dur = int(s.get("duration_sec", 0))
        total_sec += dur
        text = (s.get("narration") or "").strip()

        # 影片頁：duration 是播放長度，逐字稿只要進場與收尾的過場詞（30–120 字）
        if s.get("kind") == "video":
            video_sec += dur
            if not text:
                empty.append(str(s.get("id")))
            else:
                got = narration_chars(text)
                if got > 120:
                    r.warn(f"{s.get('id')} 是影片頁，逐字稿 {got} 字太長，"
                           "只要進場與收尾的過場詞（30–120 字）")
            continue

        if not text:
            empty.append(str(s.get("id")))
            continue
        want = dur * cpm / 60.0
        got = narration_chars(text)
        if want <= 0:
            continue
        if not (want * (1 - tol) <= got <= want * (1 + tol)):
            r.warn(f"{s.get('id')} 逐字稿 {got} 字，應為 {want:.0f} 字 ±{tol:.0%}"
                   f"（{want*(1-tol):.0f}–{want*(1+tol):.0f}）")

    if empty:
        r.fail(f"{len(empty)} 頁沒有逐字稿：{', '.join(empty[:15])}"
               + (" …" if len(empty) > 15 else ""))
        r.note("請 Claude Code 依 prompts/narration.md 寫進 deck.json 的 narration 欄位")

    m, sec = divmod(total_sec, 60)
    if not (lo_m * 60 <= total_sec <= hi_m * 60):
        r.warn(f"總時長 {m}:{sec:02d} 不在 {lo_m}–{hi_m} 分鐘區間。時間是指引："
               "超時就把次要證據頁移到附錄（role=appendix），不要刪內容；"
               "或 python tools/repace_deck.py 重新配速")
    else:
        r.note(f"總時長 {m}:{sec:02d}")
    if video_sec:
        r.note(f"其中影片播放 {format_timecode(video_sec)}"
               f"（佔 {video_sec / total_sec:.0%}）")
    return r



# --- 敘事結構（金字塔：先結論，每幕主張→證據→意涵，反方在後）------------------
def _slide_blob(s: dict, include_narration: bool = False) -> str:
    import json
    return json.dumps({k: v for k, v in s.items() if include_narration or k != "narration"},
                      ensure_ascii=False)


def check_structure(deck: dict) -> Result:
    r = Result("敘事結構", "導讀體：作者頁與執行摘要在序幕；每一主線章有章名頁籤（官方章名或原文）／"
                          "大綱頁／示意圖或今天的數字；全書意涵、反方、結語；無殘留【待填】；節奏只 WARN")
    slides = slides_of(deck)
    talk = talk_slides(slides)
    meta = deck.get("meta") or {}
    struct = meta.get("structure")
    if not struct or struct.get("form") != "guided":
        r.fail("deck.meta.structure 不是導讀體：藍圖要從 work/05a_guide.json 生成"
               "（python scripts/05_outline.py --guide-prompt）")
        return r
    if not talk:
        r.fail("沒有講述頁")
        return r

    todo = [str(s.get("id")) for s in slides if "【待填" in _slide_blob(s)]
    if todo:
        r.fail(f"{len(todo)} 頁還有【待填】：{', '.join(todo[:12])}" + (" …" if len(todo) > 12 else ""))

    roles = [s.get("role") for s in talk]
    if "author" not in roles:
        r.fail("沒有作者頁（role=author）：聽眾要先見到作者（python scripts/04_research.py --author）")
    elif roles.index("author") > 4:
        r.warn("作者頁太後面，應該緊接封面")
    if "summary" not in roles[:8]:
        r.fail("執行摘要（role=summary）要在序幕：作者頁之後、第一章之前")
    if "map" not in roles[:9]:
        r.warn("沒有章序地圖頁（role=map）")
    first_div = next((i for i, s in enumerate(talk) if s.get("role") in ("divider", "part")), None)
    if first_div is None:
        r.fail("沒有任何章名頁籤（role=divider）")

    names: set[str] = set()
    for c in meta.get("chapters") or []:
        for k in ("display", "title_en", "title"):
            if c.get(k):
                names.add(c[k].strip())
    main = struct.get("main_ch_ids") or []
    if not main:
        r.fail("meta.structure.main_ch_ids 是空的")
    n_ev = n_vis = n_today = 0
    for cid in main:
        mine = [s for s in talk if s.get("section") == cid]
        rc = Counter(s.get("role") for s in mine)
        cname = next((c.get("display") for c in meta.get("chapters") or [] if c["ch_id"] == cid), cid)
        label = f"{cid}「{cname}」"
        if not rc.get("divider"):
            r.fail(f"{label} 沒有章名頁籤")
        for s in mine:
            if s.get("role") == "divider" and (s.get("title") or "").strip() not in names:
                r.fail(f"{s.get('id')} 頁籤「{s.get('title')}」不是這本書的章名：用官方中譯或原文，不要自己改寫")
        if not rc.get("outline"):
            r.fail(f"{label} 沒有大綱頁（role=outline）")
        ev = [s for s in mine if s.get("role") in ("evidence", "today")]
        n_ev += len(ev)
        if not ev:
            r.fail(f"{label} 沒有示意圖／重點／今天的數字（role=evidence|today）")
        vis = [s for s in ev if _is_visual_slide(s)]
        n_vis += len(vis)
        if ev and not vis:
            r.warn(f"{label} 沒有示意圖：至少一頁 flow／timeline／split／table 或書中原圖")
        if rc.get("today"):
            n_today += 1
        else:
            r.warn(f"{label} 沒有「今天的數字」頁（role=today）：書寫完之後這章的數字變了嗎")

    idx_imp = [i for i, s in enumerate(talk) if s.get("role") == "implication"]
    idx_counter = [i for i, s in enumerate(talk) if s.get("role") == "counter"]
    last_ch = max((i for i, s in enumerate(talk) if s.get("section") in set(main)), default=-1)
    idx_close = next((i for i, s in enumerate(talk) if s.get("role") == "closing"), None)
    if not idx_imp:
        r.fail("沒有全書意涵頁（role=implication）：對長期投資的意義")
    elif idx_imp[0] < last_ch:
        r.warn("全書意涵頁出現在章節中間，應該在所有章之後")
    if not idx_counter:
        r.fail("沒有反方頁（role=counter）：專業的讀書分享一定要講書站不住的地方")
    else:
        if idx_counter[0] < last_ch:
            r.warn("反方頁出現在章節中間，應該在所有章之後、結語之前")
        if idx_close is not None and idx_counter[-1] > idx_close:
            r.warn("反方頁排在結語之後")
    if idx_close is None:
        r.fail("沒有結語頁（role=closing）")
    else:
        # 結語之後只允許祝賀頁（config 的 deck.closing_wish，例如「業績長紅」）；
        # 附錄排在祝賀頁後面，但不算講述頁。
        tail = [s.get("role") for s in talk[idx_close + 1:]]
        if [x for x in tail if x != "wish"]:
            r.warn(f"結語之後還有 {len(tail)} 頁講述頁，結語應該排在最後（祝賀頁除外）")
        elif tail.count("wish") > 1:
            r.warn("祝賀頁不只一張")
    app_idx = [i for i, s in enumerate(slides) if is_appendix(s)]
    close_abs = next((i for i, s in enumerate(slides) if s.get("role") == "closing"), None)
    if app_idx and close_abs is not None and min(app_idx) < close_abs:
        r.warn("附錄頁出現在結語之前，附錄應該全部排在結語之後")

    # 節奏指引（config pacing）：只 WARN
    pacing = load_project().get("pacing") or {}
    d_max = int(pacing.get("divider_max_sec", 20))
    qs_max = int(pacing.get("quote_stat_max_sec", 30))
    p_max = int(pacing.get("page_max_sec", 120))
    pro_max = float(pacing.get("prologue_share_max", 0.15))
    total = sum(int(s.get("duration_sec", 0)) for s in talk)
    for s in talk:
        dur, sid = int(s.get("duration_sec", 0)), s.get("id")
        if s.get("kind") == "divider" and dur > d_max:
            r.warn(f"{sid} 頁籤 {dur} 秒 > {d_max}：頁籤只講一句話")
        elif (s.get("quote") or s.get("stat")) and dur > qs_max:
            r.warn(f"{sid} 引言／大數字頁 {dur} 秒 > {qs_max}：那是節奏工具，不停留")
        elif s.get("kind") != "video" and dur > p_max:
            r.warn(f"{sid} 停留 {dur} 秒 > {p_max}：拆頁")
    if first_div is not None and total:
        pro_sec = sum(int(s.get("duration_sec", 0)) for s in talk[:first_div])
        if pro_sec / total > pro_max:
            r.warn(f"序幕佔 {pro_sec / total:.0%} > {pro_max:.0%}（作者頁加執行摘要），聽眾等太久才進第一章")

    if r.status == "PASS":
        r.note(f"主線 {len(main)} 章、示意圖／重點／今天 {n_ev} 頁（其中視覺 {n_vis}、今天的數字 {n_today} 章）、"
               f"反方頁 {len(idx_counter)}" + (f"、附錄 {len(app_idx)} 頁" if app_idx else ""))
    return r


def _digest_of(ch_id: str) -> dict:
    p = DIGEST / f"{ch_id}.json"
    if not p.exists():
        return {}
    try:
        return read_json(p) or {}
    except Exception:  # noqa: BLE001
        return {}


def _chapter_pages(deck: dict, ch_id: str) -> list[dict]:
    return [s for s in talk_slides(slides_of(deck)) if s.get("section") == ch_id]


# --- 原書用字（每一章的投影片要出現作者的用語）--------------------------------
def check_wording(deck: dict) -> Result:
    r = Result("原書用字", "每一主線章的投影片至少出現該章一個作者用語（digest.key_terms 的 en 或 zh）")
    main = ((deck.get("meta") or {}).get("structure") or {}).get("main_ch_ids") or []
    if not main:
        r.skip("不是導讀體藍圖")
        return r
    missing, no_terms, hits = [], [], 0
    for cid in main:
        d = _digest_of(cid)
        terms = [(kt.get("en") or "").strip().lower() for kt in (d.get("key_terms") or [])] + \
                [(kt.get("zh") or "").strip() for kt in (d.get("key_terms") or [])]
        terms = [x for x in terms if x]
        if not terms:
            no_terms.append(cid)
            continue
        blob = " ".join(_slide_blob(s) for s in _chapter_pages(deck, cid))
        low = blob.lower()
        if any(x in low for x in terms):
            hits += 1
        else:
            missing.append(cid)
    if no_terms:
        r.warn(f"{len(no_terms)} 章的 digest 沒有 key_terms（先跑 03_digest.py --supplement）：{', '.join(no_terms[:8])}")
    if missing:
        r.warn(f"{len(missing)} 章的投影片沒有出現作者用語，主標或內文改用他的話：{', '.join(missing[:10])}")
    if r.status == "PASS":
        r.note(f"{hits}/{len(main)} 章的投影片都有作者用語")
    return r


# --- 故事（每一章的逐字稿要講到選定的故事）----------------------------------
def check_stories(deck: dict) -> Result:
    r = Result("故事", "每一主線章的逐字稿要講到該章選的故事（大綱頁 story_hint 的人物）")
    main = ((deck.get("meta") or {}).get("structure") or {}).get("main_ch_ids") or []
    if not main:
        r.skip("不是導讀體藍圖")
        return r
    if not any((s.get("narration") or "").strip() for s in talk_slides(slides_of(deck))):
        r.skip("逐字稿還沒寫，之後再檢查")
        return r
    missing, no_hint = [], []
    for cid in main:
        pages = _chapter_pages(deck, cid)
        hint = next(((s.get("story_hint") or "") for s in pages if s.get("story_hint")), "")
        if not hint:
            no_hint.append(cid)
            continue
        who = re.split(r"[（(，,]", hint, 1)[0].strip()
        narr = " ".join((s.get("narration") or "") for s in pages)
        # 逐字稿是口語，「拉格什統治者恩美鐵那」講出來會變成「拉格什的統治者恩美鐵那」，
        # 所以除了整串，也接受：原文人名（story_hint 括號裡的拉丁字）、稱謂後面的名字（末四字）。
        keys = [who] + [w for w in re.split(r"[\s・·]", who) if len(w) >= 2]
        keys += re.findall(r"[A-Z][A-Za-z.\-]+(?:\s+[A-Z][A-Za-z.\-]+)+", hint)
        if len(who) >= 5:
            keys.append(who[-4:])
        if not any(k and k.lower() in narr.lower() for k in keys):
            missing.append(f"{cid}（{who}）")
    if no_hint:
        r.warn(f"{len(no_hint)} 章的大綱頁沒有 story_hint（藍圖不是最新版）：{', '.join(no_hint[:8])}")
    if missing:
        r.warn(f"{len(missing)} 章的逐字稿沒講到故事，開場用它：{', '.join(missing[:8])}")
    if r.status == "PASS":
        r.note(f"{len(main)} 章的逐字稿都講到了故事")
    return r


# --- 條列樣式（不准裸條列）----------------------------------------------------
def check_body_styles(deck: dict) -> Result:
    lim = limits()
    styles = tuple(lim.get("body_styles") or BODY_STYLES)
    label_max = int(lim.get("label_max_chars", 5))
    c_lo, c_hi = (lim.get("chain_steps_range") or [2, 4])[:2]
    p_max = int(lim.get("prose_paragraphs_max", 2))
    p_chars = int(lim.get("prose_paragraph_max_chars", 90))
    b_lo, b_hi = (lim.get("bullets_l1_range") or [2, 3])[:2]
    l1_max = int(lim.get("bullet_l1_max_chars", 40))

    r = Result("條列樣式", f"文字頁必為 {'/'.join(styles)} 之一；標籤 ≤{label_max} 字；"
                          f"鏈 {c_lo}–{c_hi} 步；敘事 ≤{p_max} 段、每段 ≤{p_chars} 字")
    counts: Counter = Counter()
    for s in slides_of(deck):
        if s.get("kind") != "content":
            continue
        body = [b for b in (s.get("body") or []) if (b.get("text") or "").strip()]
        if not body:
            continue
        sid, style = s.get("id"), s.get("style")
        if style not in styles:
            r.fail(f"{sid} 是裸條列（style={style!r}）：改成 chain（論證鏈）／"
                   "labeled（標籤＋說明）／prose（敘事段）之一，見 prompts/outline.md")
            continue
        counts[style] += 1
        if any(int(b.get("level", 0)) > 0 for b in body):
            r.fail(f"{sid} {style} 樣式不准有第二層（level>0）")
        if style == "labeled":
            if not (b_lo <= len(body) <= b_hi):
                r.warn(f"{sid} labeled {len(body)} 條，建議 {b_lo}–{b_hi} 條")
            labels = []
            for b in body:
                lab = (b.get("label") or "").strip()
                if not lab:
                    r.fail(f"{sid} labeled 樣式每一條都要有 label")
                elif visual_len(lab) > label_max:
                    r.fail(f"{sid} 標籤「{lab}」{visual_len(lab):.0f} 字 > {label_max}")
                labels.append(lab)
            if len(set(labels)) < len(labels):
                r.warn(f"{sid} 標籤重複：{labels}")
        elif style == "chain":
            if not (c_lo <= len(body) <= c_hi):
                r.fail(f"{sid} chain {len(body)} 步，必須 {c_lo}–{c_hi} 步")
        else:                                                    # prose
            if len(body) > p_max:
                r.fail(f"{sid} prose {len(body)} 段 > {p_max}")
            for b in body:
                n = visual_len(b.get("text") or "")
                if n > p_chars:
                    r.fail(f"{sid} prose 段落 {n:.0f} 字 > {p_chars}")
    total = sum(counts.values())
    if total >= 6:
        top, n = counts.most_common(1)[0]
        if n / total > 0.7:
            r.warn(f"文字頁 {n}/{total} 都是 {top}，三種樣式要混用：機制用 chain、"
                   "對照用 labeled、有故事的用 prose")
    if r.status == "PASS":
        r.note("、".join(f"{k}×{v}" for k, v in counts.items()) or "沒有文字頁")
    return r


# --- 節奏（條列頁不要連成一片）------------------------------------------------
VISUAL_KEYS = ("image", "flow", "timeline", "table", "split", "quote", "stat", "chart")


def _is_bullet_slide(s: dict) -> bool:
    """純文字條列頁：有 body、且沒有任何視覺元素。"""
    if s.get("kind") != "content":
        return False
    if any(s.get(k) for k in VISUAL_KEYS):
        return False
    return bool(s.get("body"))


def _is_visual_slide(s: dict) -> bool:
    return any(s.get(k) for k in VISUAL_KEYS)


def check_rhythm(deck: dict) -> Result:
    lim = limits()
    max_run = int(lim.get("max_consecutive_bullet_slides", 2))
    min_ratio = float(lim.get("visual_slide_ratio_min", 0.40))

    r = Result("節奏", f"連續條列頁 ≤ {max_run}；視覺頁佔比 ≥ {min_ratio:.0%}（附錄不計）")
    slides = talk_slides(slides_of(deck))
    if not slides:
        r.skip("沒有投影片")
        return r

    # 連續條列
    run, run_start = 0, None
    worst = []
    for s in slides:
        if _is_bullet_slide(s):
            run += 1
            if run == 1:
                run_start = s.get("id")
            if run == max_run + 1:
                worst.append((run_start, run))
            elif run > max_run + 1 and worst:
                worst[-1] = (run_start, run)
        else:
            run, run_start = 0, None
    for sid, n in worst:
        r.fail(f"{sid} 起連續 {n} 頁條列（上限 {max_run}），"
               "中間插一頁視覺頁、引言頁或大數字頁")

    # 視覺頁比例
    n_vis = sum(1 for s in slides if _is_visual_slide(s))
    ratio = n_vis / len(slides)
    if ratio < min_ratio:
        r.warn(f"視覺頁只佔 {ratio:.0%}（{n_vis}/{len(slides)}），低於 {min_ratio:.0%}。"
               "把「是一個機制／一組對照／一條時間線」的條列頁改畫成 "
               "flow / table / split / timeline")
    else:
        r.note(f"視覺頁 {n_vis}/{len(slides)}（{ratio:.0%}）")

    n_bullet = sum(1 for s in slides if _is_bullet_slide(s))
    r.note(f"純條列頁 {n_bullet}/{len(slides)}（{n_bullet / len(slides):.0%}）")
    return r


# --- 9. 對岸用語 ----------------------------------------------------------
def check_banned_terms(deck: dict) -> Result:
    banned = load_project().get("banned_terms", []) or []
    r = Result("對岸用語", f"黑名單掃描（{len(banned)} 個詞）")
    if not banned:
        r.skip("config/project.yaml 沒有設定 banned_terms")
        return r
    hits: Counter = Counter()
    where: dict[str, list[str]] = {}
    for s in slides_of(deck):
        blob = _slide_blob(s, include_narration=True)
        for term in banned:
            if term and term in blob:
                hits[term] += blob.count(term)
                where.setdefault(term, []).append(str(s.get("id")))
    if hits:
        for term, n in hits.most_common():
            r.fail(f"「{term}」出現 {n} 次：{', '.join(where[term][:6])}"
                   + (" …" if len(where[term]) > 6 else ""))
    else:
        r.note("沒有偵測到對岸用語")
    return r


# --- AI 味：投影片讀起來像不像人寫的 -------------------------------------
_LATIN_WORD = re.compile(r"[A-Za-z][A-Za-z'’.\-]{1,}")
# 這些字出現代表句子裡有動作或判斷，不是純名詞堆疊
_VERBISH = ("是", "不", "會", "要", "把", "被", "讓", "在", "有", "沒", "從", "到", "回",
            "漲", "跌", "升", "降", "借", "付", "買", "賣", "做", "走", "看", "說", "剩",
            "變", "算", "換", "壓", "拉", "撐", "生", "收", "花", "欠", "只", "才", "就")


def check_ai_tone(deck: dict) -> Result:
    """實跑《時間的代價》時使用者的原話：「用字不夠平易近人，太 AI 感」。

    拆解出三個可以量的特徵，全部只 WARN（語氣是人判斷的，程式只負責提醒）：
      1. 冒號句式  「名詞：名詞」一頁三條以上，整頁就像詞條表
      2. 英文夾雜  一頁超過兩個英文原詞，中文就被切碎了
      3. 排比      三條以上長度幾乎一樣，是湊出來的對仗，不是想出來的話
    另外檢查條列有沒有動詞（純名詞堆疊）。
    """
    r = Result("AI 味", "投影片不要：一頁三條冒號句式／超過兩個英文原詞／三條等長排比")
    colon_pages, latin_pages, parallel_pages, nounish = [], [], [], []
    total_items = colon_items = 0
    # 封面、章序地圖、頁籤、結語、祝賀頁的內文是人名與章名，不是句子
    skip_roles = {"cover", "map", "divider", "part", "closing", "wish"}

    for s in talk_slides(slides_of(deck)):
        if s.get("role") in skip_roles:
            continue
        sid = s.get("id")
        texts = [(b.get("text") or "").strip() for b in (s.get("body") or [])
                 if (b.get("text") or "").strip()]
        if not texts:
            continue
        total_items += len(texts)
        n_colon = sum(1 for x in texts if "：" in x or ":" in x)
        colon_items += n_colon
        if n_colon >= 3:
            colon_pages.append(sid)

        # 作者頁的人名、書名、獎項本來就是英文，不列入夾雜檢查
        if s.get("role") != "author":
            blob = " ".join(texts) + " " + (s.get("title") or "") + " " + (s.get("subtitle") or "")
            words = {w.lower() for w in _LATIN_WORD.findall(blob) if len(w) > 1}
            if len(words) > 2:
                latin_pages.append(f"{sid}（{len(words)} 個）")

        if len(texts) >= 3:
            lens = sorted(visual_len(x) for x in texts)
            if lens[-1] - lens[0] <= 3:
                parallel_pages.append(sid)

        for x in texts:
            # 有數字就當成有實質內容（「2023 年 Hayek Book Prize 得主」不算堆疊）
            if not any(v in x for v in _VERBISH) and not re.search(r"[0-9０-９]", x):
                nounish.append(f"{sid}「{x[:14]}」")

    if colon_pages:
        r.warn(f"{len(colon_pages)} 頁有三條以上冒號句式，改寫成完整的話："
               f"{', '.join(colon_pages[:8])}")
    if total_items and colon_items / total_items > 0.5:
        r.warn(f"全書 {colon_items}/{total_items} 條是冒號句式（{colon_items / total_items:.0%}），"
               "整份看起來像詞條表")
    if latin_pages:
        r.warn(f"{len(latin_pages)} 頁英文原詞超過兩個，只留作者自己造的詞："
               f"{', '.join(latin_pages[:8])}")
    if parallel_pages:
        r.warn(f"{len(parallel_pages)} 頁的要點幾乎等長，像湊出來的排比："
               f"{', '.join(parallel_pages[:8])}")
    if nounish:
        r.warn(f"{len(nounish)} 條沒有動作詞，是名詞堆疊：{', '.join(nounish[:6])}")
    if r.status == "PASS":
        r.note("沒有偵測到冒號句式、英文夾雜與等長排比")
    return r


# --- PDF：PPT 換台電腦會跑版，PDF 是交付保證 -------------------------------
def check_pdf(pptx: Path | None) -> Result:
    r = Result("PDF", "output 有同名 PDF，且頁數與 PPTX 一致")
    if not pptx or not pptx.exists():
        r.skip("找不到 PPTX")
        return r
    pdf = pptx.with_suffix(".pdf")
    if not pdf.exists():
        r.fail(f"沒有 {pdf.name}：本機請裝 libreoffice-impress 後重跑 06_build_pptx.py")
        return r
    try:
        from pptx import Presentation
        n_ppt = len(Presentation(str(pptx)).slides)
    except Exception:
        n_ppt = 0
    n_pdf = 0
    try:
        import fitz
        n_pdf = fitz.open(str(pdf)).page_count
    except Exception:
        data = pdf.read_bytes()
        n_pdf = data.count(b"/Type /Page") - data.count(b"/Type /Pages")
    if n_ppt and n_pdf and n_ppt != n_pdf:
        r.fail(f"PDF {n_pdf} 頁 ≠ PPTX {n_ppt} 頁：PDF 是舊的，請重跑 06_build_pptx.py")
    else:
        r.note(f"{pdf.name}（{n_pdf or '?'} 頁）")
    return r


# --- 10. 視覺預覽 ---------------------------------------------------------
def make_preview(pptx: Path | None) -> Result:
    r = Result("視覺預覽", "LibreOffice 轉 PNG 全頁截圖 → output/preview/")
    if pptx is None or not pptx.exists():
        r.skip("找不到 output/*.pptx")
        return r

    so = soffice_bin()
    if not so:
        r.skip("找不到 LibreOffice。Ubuntu: sudo apt install libreoffice-impress ／ "
               "macOS: brew install --cask libreoffice")
        return r

    outdir = OUTPUT / "preview"
    outdir.mkdir(parents=True, exist_ok=True)
    pdf = OUTPUT / (pptx.stem + ".pdf")
    try:
        proc = subprocess.run(
            [so, "--headless", "--convert-to", "pdf", str(pptx), "--outdir", str(OUTPUT)],
            capture_output=True, text=True, timeout=600,
        )
    except subprocess.TimeoutExpired:
        r.skip("LibreOffice 轉檔逾時（600s）")
        return r

    if not pdf.exists():
        msg = (proc.stderr or proc.stdout or "").strip().splitlines()
        detail = msg[-1] if msg else "無錯誤訊息"
        r.skip(f"LibreOffice 無法轉檔（{detail}）。"
               "多半是只裝了 libreoffice-core、缺 Impress 匯入濾鏡：\n"
               "        sudo apt install libreoffice-impress")
        return r

    # 優先用 pdftoppm；沒有就用 pymupdf
    n = 0
    if have_cmd("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", "80", str(pdf), str(outdir / "slide")],
                       capture_output=True, timeout=600)
        n = len(list(outdir.glob("slide*.png")))
    else:
        try:
            import pymupdf

            doc = pymupdf.open(pdf)
            for i, page in enumerate(doc, start=1):
                page.get_pixmap(dpi=80).save(str(outdir / f"slide-{i:03d}.png"))
            n = len(doc)
            doc.close()
            r.note("（pdftoppm 不存在，改用 pymupdf 產圖）")
        except Exception as e:  # noqa: BLE001
            r.skip(f"PDF 轉 PNG 失敗：{e}")
            return r

    r.note(f"{n} 張預覽圖 → {outdir}／PDF → {pdf}")
    return r


# ==========================================================================
def check_env() -> int:
    step("環境檢查")
    cfg_ok = True

    print("  Python 套件")
    for mod, why in [("pptx", "產 PPTX"), ("docx", "產逐字稿"), ("pymupdf", "解析 PDF"),
                     ("ebooklib", "解析 EPUB"), ("bs4", "EPUB HTML"), ("yaml", "讀設定"),
                     ("matplotlib", "畫圖表"), ("PIL", "圖片處理"),
                     ("rapidfuzz", "頁眉偵測"), ("requests", "連結檢查"),
                     ("pytesseract", "OCR（掃描版 PDF 才需要）")]:
        try:
            __import__(mod)
            print(_c("32", f"    ✓ {mod:<12} {why}"))
        except ImportError:
            optional = mod == "pytesseract"
            print(_c("33" if optional else "31",
                     f"    {'!' if optional else '✗'} {mod:<12} {why} — 未安裝"))
            if not optional:
                cfg_ok = False

    print("\n  外部工具")
    for cmd, why, optional in [
        ("tesseract", "OCR 掃描版 PDF", True),
        ("soffice", "QA 預覽截圖", True),
        ("pdftoppm", "PDF 轉 PNG（沒有會改用 pymupdf）", True),
    ]:
        if have_cmd(cmd):
            print(_c("32", f"    ✓ {cmd:<12} {why}"))
        else:
            print(_c("33", f"    ! {cmd:<12} {why} — 未安裝（該項檢查會標 SKIP）"))

    # soffice 存在不等於能轉檔（只裝 core、沒裝 impress 的情況）
    if soffice_bin():
        print("\n  LibreOffice 轉檔能力")
        probe = _probe_soffice()
        print(_c("32" if probe else "33",
                 f"    {'✓' if probe else '!'} "
                 + ("可以轉檔" if probe else
                    "soffice 在但無法轉檔（多半缺 libreoffice-impress）→ 預覽會標 SKIP")))

    print("\n  資產")
    for p, why, required in [(TEMPLATE, "母片", True),
                             (Path("config/fh_template_spec.json"), "版面規格", True),
                             (Path("config/project.yaml"), "專案設定", True)]:
        full = Path(p) if Path(p).is_absolute() else Path(__file__).resolve().parent.parent / p
        if full.exists():
            print(_c("32", f"    ✓ {why:<10} {full.name}（{full.stat().st_size / 1024:.0f} KB）"))
        else:
            print(_c("31" if required else "33", f"    ✗ {why:<10} 找不到 {full}"))
            if required:
                cfg_ok = False

    print("\n  中文字型（影響圖表與 QA 預覽，不影響 PPTX 本身）")
    fonts = _cjk_fonts()
    if fonts:
        for f in fonts[:4]:
            print(_c("32", f"    ✓ {f}"))
    else:
        print(_c("33", "    ! 找不到中文字型 — 圖表中文會變方框。"
                       "Ubuntu: sudo apt install fonts-noto-cjk"))

    print()
    if cfg_ok:
        ok("必要相依齊全，可以開跑")
        info("下一步：把書檔放進 input/，然後 make extract")
        return 0
    print(_c("31", "  ✗ 有缺少的必要相依，請跑 make setup"))
    return 1


def _probe_soffice() -> bool:
    """實際做一次最小轉檔，確認匯入／匯出濾鏡真的在。"""
    import tempfile

    so = soffice_bin()
    if not so:
        return False
    try:
        from pptx import Presentation
    except ImportError:
        return False
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / "probe.pptx"
        Presentation().save(str(src))
        try:
            subprocess.run([so, "--headless", "--convert-to", "pdf", str(src),
                            "--outdir", td], capture_output=True, timeout=180)
        except subprocess.TimeoutExpired:
            return False
        return (Path(td) / "probe.pdf").exists()


def _cjk_fonts() -> list[str]:
    try:
        out = subprocess.run(["fc-list", ":lang=zh-tw", "family"],
                             capture_output=True, text=True, timeout=20).stdout
    except Exception:  # noqa: BLE001
        return []
    names = []
    for line in out.splitlines():
        for part in line.split(","):
            part = part.strip()
            if part and part not in names:
                names.append(part)
    return names


# ==========================================================================
def write_report(results: list[Result], deck: dict, pptx: Path | None) -> Path:
    meta = deck.get("meta", {})
    n_fail = sum(1 for r in results if r.status == "FAIL")
    verdict = "**不可交付**" if n_fail else "**可以交付**"

    lines = [
        "# QA 報告",
        "",
        f"- 產生時間：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- 書名：{meta.get('book_title', '')}",
        f"- 講者：{meta.get('dept', '')} {meta.get('presenter', '')}",
        f"- 投影片：`{pptx}`" if pptx else "- 投影片：（尚未產出）",
        f"- 藍圖：`{DECK_JSON}`（{len(deck.get('slides', []))} 頁）",
        "",
        f"## 結論：{verdict}",
        "",
        "| 檢查 | 判準 | 結果 |",
        "|---|---|---|",
    ]
    for r in results:
        lines.append(f"| {r.name} | {r.criterion} | **{r.status}** |")
    lines += ["", "> 任何一項 FAIL 就不准交付（規劃書 §10）。", ""]

    for r in results:
        if not r.details:
            continue
        lines += [f"## {r.name} — {r.status}", "", f"判準：{r.criterion}", ""]
        for d in r.details:
            lines.append(f"- {d}")
        lines.append("")

    lines += [
        "## 修正流程",
        "",
        "1. 改 `work/05_deck.json`（文案、頁數、duration_sec、sources）",
        "2. `make revise`（重跑 Stage 6–8，不要重跑前面的階段）",
        "3. 回到本報告確認 FAIL 歸零",
        "",
    ]
    out = OUTPUT / "qa_report.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


if __name__ == "__main__":
    sys.exit(main())
