#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""從 work/05_deck.json 匯出「該去搜什麼圖」的清單。

使用者的需求：簡報裡建議放哪些圖、用什麼關鍵字去搜，他自己抓回來加工。
圖檔不在這裡抓——雲端的對外連線是白名單，機構網站與圖庫幾乎全被擋。

    python tools/image_suggestions.py              # → output/圖片建議.md
    python tools/image_suggestions.py --docx       # 順便出一份 Word

每一筆的欄位（deck.json 每一頁的 image_hint）：

    what          這張圖要拍到什麼（人物／地點／年代）
    keywords_zh   中文搜尋關鍵字，直接貼進搜尋框
    keywords_en   英文搜尋關鍵字
    source        建議去哪裡找（授權要自己確認）
    use           這張圖放在哪、要做什麼
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from _common import (  # noqa: E402
    DECK_JSON, OUTPUT, die, info, ok, read_json, step, warn,
)

HEADER = """# 圖片建議清單

這份是「該去搜什麼圖」，不是圖檔本身。抓回來之後放進投影片，
把章名頁籤上那個虛線佔位框刪掉即可。

**授權要自己確認**：Wikimedia Commons 要看每張圖的授權標示；
機構官網的圖多半只能引用不能改；圖庫（Unsplash、Pexels）可商用但要看條款。

| 頁 | 頁面 | 要拍到什麼 | 中文關鍵字 | 英文關鍵字 | 去哪找 |
|---|---|---|---|---|---|
"""


def rows(deck: dict) -> list[dict]:
    out = []
    for i, s in enumerate(deck.get("slides") or [], start=1):
        hint = s.get("image_hint")
        if not hint:
            continue
        out.append({
            "page": i,
            "id": s.get("id", ""),
            "title": (s.get("title") or "").strip(),
            "role": s.get("role", ""),
            "what": (hint.get("what") or "").strip(),
            "zh": (hint.get("keywords_zh") or "").strip(),
            "en": (hint.get("keywords_en") or "").strip(),
            "source": (hint.get("source") or "").strip(),
            "use": (hint.get("use") or "").strip(),
        })
    return out


def write_md(items: list[dict], path: Path) -> None:
    lines = [HEADER]
    for r in items:
        lines.append(f"| {r['page']} | {r['title']} | {r['what']} | {r['zh']} | "
                     f"{r['en']} | {r['source']} |\n")
    lines.append("\n## 逐頁說明\n\n")
    for r in items:
        lines.append(f"### 第 {r['page']} 頁　{r['title']}\n\n")
        lines.append(f"- **要拍到什麼**：{r['what']}\n")
        lines.append(f"- **中文關鍵字**：`{r['zh']}`\n")
        if r["en"]:
            lines.append(f"- **英文關鍵字**：`{r['en']}`\n")
        lines.append(f"- **去哪找**：{r['source']}\n")
        lines.append(f"- **放哪裡**：{r['use']}\n\n")
    path.write_text("".join(lines), encoding="utf-8")


def write_docx(items: list[dict], path: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        warn("沒有 python-docx，略過 Word（pip install python-docx）")
        return False
    doc = Document()
    doc.add_heading("圖片建議清單", level=0)
    doc.add_paragraph("這份是「該去搜什麼圖」，不是圖檔本身。授權要自己確認。")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for c, txt in zip(table.rows[0].cells, ("頁", "頁面", "要拍到什麼", "中文關鍵字", "英文關鍵字")):
        c.text = txt
    for r in items:
        cells = table.add_row().cells
        for c, txt in zip(cells, (str(r["page"]), r["title"], r["what"], r["zh"], r["en"])):
            c.text = txt
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.save(str(path))
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description="匯出圖片建議清單")
    ap.add_argument("--deck", default=str(DECK_JSON))
    ap.add_argument("--docx", action="store_true", help="順便出一份 Word")
    args = ap.parse_args()

    deck_path = Path(args.deck)
    if not deck_path.exists():
        die(f"找不到 {deck_path}，請先跑 05_outline.py")
    deck = read_json(deck_path)

    step("圖片建議")
    items = rows(deck)
    if not items:
        warn("deck.json 裡沒有任何 image_hint。"
             "章名頁籤的建議由 05_outline.py 自動產生；"
             "其他頁想加圖請依 prompts/visuals.md 手動補 image_hint")
        return 0

    OUTPUT.mkdir(parents=True, exist_ok=True)
    md = OUTPUT / "圖片建議.md"
    write_md(items, md)
    ok(f"{len(items)} 筆 → {md}")
    if args.docx and write_docx(items, OUTPUT / "圖片建議.docx"):
        ok(f"Word → {OUTPUT / '圖片建議.docx'}")
    info("抓回來之後放進投影片，把章名頁籤上的虛線佔位框刪掉")
    return 0


if __name__ == "__main__":
    sys.exit(main())
